import io
import zipfile

from openpyxl import load_workbook
from PIL import Image

import catalogue_excel as base
import catalogue_excel_final as final


def _tiny_png():
    out = io.BytesIO()
    Image.new("RGB", (20, 12), "white").save(out, format="PNG")
    return out.getvalue()


def test_real_product_image_hosts_are_allowed_for_excel_export():
    assert "customer-assets.emergentagent.com" in base._ALLOWED_EXTERNAL_IMAGE_HOSTS
    assert "d3adwkbyhxyrtq.cloudfront.net" in base._ALLOWED_EXTERNAL_IMAGE_HOSTS
    assert "d33sy5i8bnduwe.cloudfront.net" in base._ALLOWED_EXTERNAL_IMAGE_HOSTS


def test_compatible_workbook_removes_table_parts_but_keeps_filter_and_image():
    products = [
        {
            "id": "p-1",
            "sku": "SGE-CH-001",
            "name": "Test Chandelier",
            "category": "Chandelier",
            "status": "published",
            "price_display": "starting_from",
            "price": 12500,
            "currency": "INR",
            "images": ["https://customer-assets.emergentagent.com/test.jpg"],
            "specs": {"Lights": "8"},
        }
    ]

    payload, metadata = final._build_compatible_workbook(
        products,
        {"https://customer-assets.emergentagent.com/test.jpg": _tiny_png()},
        {"failures": {}},
        "Chandelier",
    )

    assert metadata["embedded_images"] == 1

    with zipfile.ZipFile(io.BytesIO(payload), "r") as archive:
        names = set(archive.namelist())
        assert not any(name.startswith("xl/tables/") for name in names)
        assert "xl/worksheets/sheet1.xml" in names
        sheet_xml = archive.read("xl/worksheets/sheet1.xml")
        assert b"tableParts" not in sheet_xml
        assert b"autoFilter" in sheet_xml

    wb = load_workbook(io.BytesIO(payload))
    ws = wb["Products"]
    assert ws.auto_filter.ref
    assert len(ws._images) == 1
    assert ws.tables == {}
    assert wb["Summary"]["A1"].value == "Samrat Glass Emporium — Chandelier Catalogue Export"


def test_failure_sheet_lists_failed_sku_url_and_reason():
    failed_url = "https://customer-assets.emergentagent.com/fail.jpg"
    products = [
        {
            "id": "p-2",
            "sku": "SGE-HL-001",
            "name": "Test Hanging Light",
            "category": "Hanging Light",
            "status": "published",
            "images": [failed_url],
            "specs": {},
        }
    ]

    payload, metadata = final._build_compatible_workbook(
        products,
        {},
        {"failures": {failed_url: "fetch_timeout"}},
        "Hanging Light",
    )

    assert metadata["embedded_images"] == 0
    assert metadata["image_failures"] == 1
    wb = load_workbook(io.BytesIO(payload))
    assert "Image Failures" in wb.sheetnames
    ws = wb["Image Failures"]
    assert ws["A2"].value == "SGE-HL-001"
    assert ws["B2"].value == "Test Hanging Light"
    assert ws["C2"].value == "Hanging Light"
    assert ws["D2"].value == failed_url
    assert ws["E2"].value == "fetch_timeout"


def test_strip_table_parts_is_idempotent_for_plain_zip_workbook():
    products = [
        {
            "id": "p-2",
            "sku": "SGE-TL-001",
            "name": "Test Lamp",
            "category": "Table Lamp",
            "status": "published",
            "images": [],
            "specs": {},
        }
    ]
    payload, _ = base.build_catalogue_workbook(products)
    once = final._strip_table_parts(payload)
    twice = final._strip_table_parts(once)

    wb = load_workbook(io.BytesIO(twice))
    assert wb["Products"].max_row == 2
    assert wb["Products"].tables == {}
