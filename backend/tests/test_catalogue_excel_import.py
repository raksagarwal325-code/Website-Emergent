import io

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from catalogue_excel import build_catalogue_workbook
from catalogue_excel_import import (
    CatalogueImportError,
    build_import_plan,
    history_admin,
    parse_catalogue_workbook,
    selected_import_plan,
)


def _export(products):
    payload, _ = build_catalogue_workbook(products)
    return payload


def _products():
    return [
        {
            "id": "p-1",
            "sku": "SGE-HL-001",
            "name": "Kandil Clear Hanging Light",
            "category": "Hanging Light",
            "status": "published",
            "price": 18000,
            "images": ["/api/files/one.jpg"],
            "specs": {"Glass Colour": "Clear", "Glass Cut / Design": "Classic"},
        },
        {
            "id": "p-2",
            "sku": "SGE-HL-002",
            "name": "Kandil Ruby Hanging Light",
            "category": "Hanging Light",
            "status": "draft",
            "price": 19000,
            "specs": {"Glass Colour": "Ruby Red"},
        },
    ]


def _edit(payload, edits, *, hide_row=None):
    workbook = load_workbook(io.BytesIO(payload))
    sheet = workbook["Products"]
    headers = [cell.value for cell in sheet[1]]
    sku_col = headers.index("SKU") + 1
    rows = {
        sheet.cell(row=row, column=sku_col).value: row
        for row in range(2, sheet.max_row + 1)
    }
    for sku, header, value in edits:
        sheet.cell(row=rows[sku], column=headers.index(header) + 1).value = value
    if hide_row:
        sheet.row_dimensions[rows[hide_row]].hidden = True
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_import_reads_hidden_rows_and_only_plans_non_empty_spec_changes():
    products = _products()
    payload = _edit(
        _export(products),
        [
            ("SGE-HL-001", "Spec: Glass Cut / Design", "Dragon Crest"),
            ("SGE-HL-001", "Product Name", "Do not import this name"),
            ("SGE-HL-002", "Spec: Glass Cut / Design", ""),
        ],
        hide_row="SGE-HL-001",
    )

    rows = parse_catalogue_workbook(payload)
    plan = build_import_plan(rows, products)

    assert plan["workbook_rows"] == 2
    assert plan["matched_count"] == 2
    assert plan["changed_product_count"] == 1
    assert plan["changed_field_count"] == 1
    assert plan["unchanged_count"] == 1
    assert plan["errors"] == []
    assert plan["can_apply"] is True
    assert plan["items"][0]["name"] == "Kandil Clear Hanging Light"
    assert plan["items"][0]["changes"] == [{
        "field": "Spec: Glass Cut / Design",
        "old": "Classic",
        "new": "Dragon Crest",
    }]
    assert plan["items"][0]["specs"]["Glass Colour"] == "Clear"


def test_import_rejects_changed_sku_via_product_id_cross_check():
    products = _products()
    payload = _edit(
        _export(products),
        [("SGE-HL-001", "Product ID", "p-2")],
    )

    plan = build_import_plan(parse_catalogue_workbook(payload), products)

    assert plan["can_apply"] is False
    assert plan["errors"][0]["message"].startswith("SKU and Product ID do not match")


def test_import_reports_unknown_sku_and_duplicate_workbook_sku():
    products = _products()
    unknown = _edit(_export(products), [("SGE-HL-001", "SKU", "SGE-HL-999")])
    unknown_plan = build_import_plan(parse_catalogue_workbook(unknown), products)
    assert unknown_plan["errors"][0]["message"] == "SKU does not exist in Admin."

    workbook = load_workbook(io.BytesIO(_export(products)))
    sheet = workbook["Products"]
    headers = [cell.value for cell in sheet[1]]
    sheet.cell(row=3, column=headers.index("SKU") + 1).value = "SGE-HL-001"
    output = io.BytesIO()
    workbook.save(output)
    with pytest.raises(CatalogueImportError, match="appears more than once"):
        parse_catalogue_workbook(output.getvalue())


def test_import_rejects_formula_in_editable_spec_cell():
    products = _products()
    payload = _edit(
        _export(products),
        [("SGE-HL-001", "Spec: Glass Cut / Design", '=CONCAT("Dragon"," Crest")')],
    )

    with pytest.raises(CatalogueImportError, match="contains a formula"):
        parse_catalogue_workbook(payload)


def test_preview_token_changes_when_admin_value_changes_after_preview():
    products = _products()
    payload = _edit(
        _export(products),
        [("SGE-HL-001", "Spec: Glass Cut / Design", "Dragon Crest")],
    )
    rows = parse_catalogue_workbook(payload)
    first = build_import_plan(rows, products)
    products[0]["specs"]["Glass Cut / Design"] = "Diamond Cut"
    second = build_import_plan(rows, products)

    assert first["preview_token"] != second["preview_token"]


def test_auth_mapping_is_normalized_for_product_history_helpers():
    admin = history_admin({"user_id": "user-1", "email": "owner@example.com"})

    assert admin.user_id == "user-1"
    assert admin.email == "owner@example.com"


def test_selected_import_plan_applies_only_reviewed_fields():
    products = _products()
    payload = _edit(
        _export(products),
        [
            ("SGE-HL-001", "Spec: Glass Colour", "Emerald Green"),
            ("SGE-HL-001", "Spec: Glass Cut / Design", "Dragon Crest"),
        ],
    )
    plan = build_import_plan(parse_catalogue_workbook(payload), products)

    selected = selected_import_plan(plan, '[{"id":"p-1","fields":["Spec: Glass Colour"]}]')

    assert selected[0]["changes"] == [{
        "field": "Spec: Glass Colour",
        "old": "Clear",
        "new": "Emerald Green",
    }]
    assert selected[0]["specs"] == {
        "Glass Colour": "Emerald Green",
        "Glass Cut / Design": "Classic",
    }


def test_selected_import_plan_rejects_unpreviewed_fields():
    products = _products()
    payload = _edit(_export(products), [("SGE-HL-001", "Spec: Glass Colour", "Emerald Green")])
    plan = build_import_plan(parse_catalogue_workbook(payload), products)

    with pytest.raises(HTTPException, match="not present in the current preview"):
        selected_import_plan(plan, '[{"id":"p-1","fields":["Spec: Price"]}]')
