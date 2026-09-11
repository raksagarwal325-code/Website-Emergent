import io
from datetime import datetime, timezone

from openpyxl import load_workbook
from PIL import Image

from catalogue_excel import build_catalogue_workbook


def _tiny_png():
    out = io.BytesIO()
    Image.new("RGB", (20, 12), "white").save(out, format="PNG")
    return out.getvalue()


def _headers(ws):
    return [cell.value for cell in ws[1]]


def test_full_catalogue_export_keeps_all_statuses_raw_specs_and_embeds_image():
    products = [
        {
            "id": "draft-1",
            "sku": "SGE-TL-002",
            "name": "Draft Table Lamp",
            "category": "Table Lamp",
            "status": "draft",
            "price_display": "on_request",
            "price": 0,
            "currency": "INR",
            "images": [],
            "tags": ["internal-tag", "collection:test"],
            "specs": {"Holder": "B22 or compatible holders; confirm before order."},
            "description": "Awaiting owner review.",
        },
        {
            "id": "published-1",
            "sku": "SGE-CH-001",
            "name": "Published Chandelier",
            "category": "Chandelier",
            "status": "published",
            "price_display": "starting_from",
            "price": 12500,
            "currency": "INR",
            "images": ["/api/files/catalog-app/products/example.jpg"],
            "tags": ["heritage", "firozabad"],
            "specs": {"Lights": "8", "Material": "Cut glass"},
            "short_description": "A catalogue chandelier.",
        },
    ]

    payload, metadata = build_catalogue_workbook(
        products,
        image_loader=lambda _url: _tiny_png(),
        generated_at=datetime(2026, 9, 11, 12, 0, tzinfo=timezone.utc),
    )

    assert payload[:2] == b"PK"  # xlsx is a ZIP container
    assert metadata == {
        "total": 2,
        "embedded_images": 1,
        "image_failures": 0,
        "no_image": 1,
    }

    wb = load_workbook(io.BytesIO(payload))
    assert wb.sheetnames == ["Products", "Summary"]
    ws = wb["Products"]
    headers = _headers(ws)
    assert "Image" in headers
    assert "Spec: Holder" in headers
    assert "Spec: Lights" in headers
    assert "Spec: Material" in headers
    assert "All Image URLs" in headers
    assert "Status" in headers
    assert len(ws._images) == 1

    by_sku = {
        ws.cell(row=row, column=headers.index("SKU") + 1).value: row
        for row in range(2, ws.max_row + 1)
    }
    assert set(by_sku) == {"SGE-CH-001", "SGE-TL-002"}

    draft_row = by_sku["SGE-TL-002"]
    assert ws.cell(draft_row, headers.index("Status") + 1).value == "draft"
    assert ws.cell(draft_row, headers.index("Price Display") + 1).value == "Price on Request"
    assert "collection:test" in ws.cell(draft_row, headers.index("Tags") + 1).value
    assert (
        ws.cell(draft_row, headers.index("Spec: Holder") + 1).value
        == "B22 or compatible holders; confirm before order."
    )

    published_row = by_sku["SGE-CH-001"]
    assert ws.cell(published_row, headers.index("Status") + 1).value == "published"
    assert ws.cell(published_row, headers.index("Primary Image URL") + 1).value.endswith("example.jpg")

    summary = wb["Summary"]
    assert summary["A3"].value == "Total products"
    assert summary["B3"].value == 2
    assert summary["A4"].value == "Embedded primary images"
    assert summary["B4"].value == 1


def test_image_failure_does_not_abort_excel_export():
    products = [
        {
            "id": "p-1",
            "sku": "SGE-WL-001",
            "name": "Wall Light",
            "category": "Wall Light",
            "status": "published",
            "images": ["https://images.unsplash.com/fail.jpg"],
            "specs": {},
        }
    ]

    payload, metadata = build_catalogue_workbook(products, image_loader=lambda _url: None)
    wb = load_workbook(io.BytesIO(payload))

    assert wb["Products"].max_row == 2
    assert metadata["total"] == 1
    assert metadata["embedded_images"] == 0
    assert metadata["image_failures"] == 1
    assert metadata["no_image"] == 0
