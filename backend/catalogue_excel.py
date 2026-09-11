"""Admin-only Excel catalogue export with embedded product thumbnails.

This module is intentionally isolated from ``server.py`` so the large API file
stays unchanged. ``auth.py`` installs the route while the FastAPI app is being
initialised. The export is an operational/admin artefact: it contains every
product row (published, draft, or other workflow state), preserves raw specs
and tags, embeds the primary product image when it can be resolved safely, and
keeps every image URL in the workbook for auditability.
"""
from __future__ import annotations

import asyncio
import io
import ipaddress
import json
import re
import socket
import sys
from collections import Counter
from datetime import datetime, timezone
from urllib.parse import unquote, urlparse

import requests
from fastapi import HTTPException, Request
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from PIL import Image as PILImage, ImageOps
from starlette.responses import Response

_SITE_ORIGIN = "https://samratglass.com"
_MAX_IMAGE_BYTES = 15 * 1024 * 1024
_ALLOWED_EXTERNAL_IMAGE_HOSTS = {
    "samratglass.com",
    "www.samratglass.com",
    "images.unsplash.com",
    "images.pexels.com",
    "images.pexelsusercontent.com",
}


def _find_server_module():
    for name in ("server", "backend.server"):
        mod = sys.modules.get(name)
        if mod is not None and hasattr(mod, "app") and hasattr(mod, "db"):
            return mod
    for mod in tuple(sys.modules.values()):
        if mod is None or not hasattr(mod, "app") or not hasattr(mod, "db"):
            continue
        path = str(getattr(mod, "__file__", "") or "")
        if path.endswith("/server.py") or path.endswith("\\server.py"):
            return mod
    return None


def _clean_cell(value, limit: int = 32000):
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple, set)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value))
    return text[:limit]


def _human_price_display(value: str) -> str:
    return {
        "on_request": "Price on Request",
        "fixed": "Fixed Price",
        "starting_from": "Starting From",
    }.get(str(value or "").strip(), str(value or "").strip() or "Starting From")


def _storage_path_from_url(raw: str) -> str | None:
    value = str(raw or "").strip()
    if not value:
        return None
    if value.startswith("/api/files/"):
        return unquote(value[len("/api/files/"):])
    if value.startswith("api/files/"):
        return unquote(value[len("api/files/"):])
    try:
        parsed = urlparse(value)
    except Exception:
        return None
    host = (parsed.hostname or "").lower().rstrip(".")
    if host in {"samratglass.com", "www.samratglass.com"} and parsed.path.startswith("/api/files/"):
        return unquote(parsed.path[len("/api/files/"):])
    return None


def _is_public_ip(raw: str) -> bool:
    try:
        addr = ipaddress.ip_address(raw)
    except ValueError:
        return False
    return not (
        addr.is_private
        or addr.is_loopback
        or addr.is_link_local
        or addr.is_multicast
        or addr.is_reserved
        or addr.is_unspecified
    )


def _safe_external_image_bytes(raw: str) -> bytes | None:
    """Fetch only known public HTTPS image hosts; never follow redirects."""
    try:
        parsed = urlparse(str(raw or "").strip())
    except Exception:
        return None
    if parsed.scheme.lower() != "https":
        return None
    host = (parsed.hostname or "").lower().rstrip(".")
    if not host or not (
        host in _ALLOWED_EXTERNAL_IMAGE_HOSTS
        or any(host.endswith("." + allowed) for allowed in _ALLOWED_EXTERNAL_IMAGE_HOSTS)
    ):
        return None
    try:
        infos = socket.getaddrinfo(host, 443, proto=socket.IPPROTO_TCP)
    except OSError:
        return None
    if not infos or any(not _is_public_ip(sockaddr[0]) for *_prefix, sockaddr in infos):
        return None
    try:
        with requests.get(raw, timeout=8, stream=True, allow_redirects=False) as response:
            if response.status_code < 200 or response.status_code >= 300:
                return None
            content_type = (response.headers.get("Content-Type") or "").split(";", 1)[0].strip().lower()
            if not content_type.startswith("image/"):
                return None
            content_length = response.headers.get("Content-Length")
            if content_length:
                try:
                    if int(content_length) > _MAX_IMAGE_BYTES:
                        return None
                except ValueError:
                    pass
            chunks: list[bytes] = []
            total = 0
            for chunk in response.iter_content(64 * 1024):
                if not chunk:
                    continue
                total += len(chunk)
                if total > _MAX_IMAGE_BYTES:
                    return None
                chunks.append(chunk)
            return b"".join(chunks)
    except requests.RequestException:
        return None


def _primary_image_bytes(raw: str, get_object) -> bytes | None:
    path = _storage_path_from_url(raw)
    if path:
        try:
            data, content_type = get_object(path)
            if data and str(content_type or "").lower().startswith("image/"):
                return data
        except Exception:
            return None
    return _safe_external_image_bytes(raw)


def _normalise_thumbnail(raw: bytes, max_side: int = 128) -> bytes | None:
    try:
        with PILImage.open(io.BytesIO(raw)) as opened:
            if getattr(opened, "is_animated", False):
                opened.seek(0)
            image = ImageOps.exif_transpose(opened)
            image.thumbnail((max_side, max_side), PILImage.Resampling.LANCZOS)
            if image.mode not in {"RGB", "RGBA"}:
                image = image.convert("RGBA" if "A" in image.getbands() else "RGB")
            out = io.BytesIO()
            image.save(out, format="PNG", optimize=True)
            return out.getvalue()
    except Exception:
        return None


def _spec_keys(products: list[dict]) -> list[str]:
    keys: set[str] = set()
    for product in products:
        specs = product.get("specs") or {}
        if isinstance(specs, dict):
            keys.update(str(k).strip() for k in specs if str(k).strip())
    return sorted(keys, key=str.casefold)


def _product_url(product: dict) -> str:
    identity = str(product.get("id") or "").strip()
    return f"{_SITE_ORIGIN}/product/{identity}" if identity else ""


def build_catalogue_workbook(
    products: list[dict],
    *,
    image_loader=None,
    generated_at: datetime | None = None,
) -> tuple[bytes, dict]:
    """Return XLSX bytes plus export metadata.

    ``image_loader`` receives the product's primary image URL and returns raw
    image bytes or ``None``. Keeping it injectable makes the workbook logic
    deterministic and easy to unit-test without network/storage access.
    """
    generated_at = generated_at or datetime.now(timezone.utc)
    ordered = sorted(
        [dict(p or {}) for p in products],
        key=lambda p: (
            str(p.get("category") or "").casefold(),
            str(p.get("sku") or "").casefold(),
            str(p.get("name") or "").casefold(),
        ),
    )
    spec_keys = _spec_keys(ordered)

    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "B2"

    base_headers = [
        "Image",
        "SKU",
        "Product Name",
        "Category",
        "Status",
        "Price Display",
        "Price",
        "Compare At Price",
        "Currency",
        "Stock",
        "Featured",
        "Badge",
        "Short Description",
        "Description",
    ]
    spec_headers = [f"Spec: {key}" for key in spec_keys]
    tail_headers = [
        "Tags",
        "Primary Image URL",
        "All Image URLs",
        "Product URL",
        "Product ID",
        "Created At",
        "Updated At",
    ]
    headers = base_headers + spec_headers + tail_headers
    ws.append(headers)

    gold = "D4AF37"
    dark = "171717"
    pale = "F6F1E3"
    thin_gold = Side(style="thin", color="D9C88B")
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=dark)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=gold))
    ws.row_dimensions[1].height = 30

    embedded = 0
    image_failures = 0
    no_image = 0
    status_counts: Counter[str] = Counter()
    category_counts: Counter[str] = Counter()

    for product in ordered:
        images = [str(x).strip() for x in (product.get("images") or []) if str(x).strip()]
        primary = images[0] if images else ""
        status = str(product.get("status") or "published")
        category = str(product.get("category") or "")
        status_counts[status] += 1
        category_counts[category or "(blank)"] += 1
        specs = product.get("specs") if isinstance(product.get("specs"), dict) else {}
        price = product.get("price")
        compare = product.get("compare_at_price")
        values = [
            "",
            _clean_cell(product.get("sku")),
            _clean_cell(product.get("name")),
            _clean_cell(category),
            _clean_cell(status),
            _clean_cell(_human_price_display(product.get("price_display"))),
            price if isinstance(price, (int, float)) else _clean_cell(price),
            compare if isinstance(compare, (int, float)) else _clean_cell(compare),
            _clean_cell(product.get("currency") or "INR"),
            product.get("stock") if isinstance(product.get("stock"), (int, float)) else _clean_cell(product.get("stock")),
            "Yes" if product.get("featured") else "No",
            _clean_cell(product.get("badge")),
            _clean_cell(product.get("short_description")),
            _clean_cell(product.get("description")),
        ]
        values.extend(_clean_cell(specs.get(key)) for key in spec_keys)
        values.extend([
            _clean_cell(", ".join(str(x) for x in (product.get("tags") or []))),
            _clean_cell(primary),
            _clean_cell("\n".join(images)),
            _product_url(product),
            _clean_cell(product.get("id")),
            _clean_cell(product.get("created_at")),
            _clean_cell(product.get("updated_at")),
        ])
        ws.append(values)
        row = ws.max_row
        ws.row_dimensions[row].height = 84

        for cell in ws[row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin_gold)
        if row % 2 == 0:
            for cell in ws[row]:
                cell.fill = PatternFill("solid", fgColor=pale)

        # Currency formatting for numeric price cells.
        for header in ("Price", "Compare At Price"):
            col = headers.index(header) + 1
            ws.cell(row=row, column=col).number_format = '₹#,##0.00;[Red]-₹#,##0.00'

        url_col = headers.index("Product URL") + 1
        if ws.cell(row=row, column=url_col).value:
            ws.cell(row=row, column=url_col).hyperlink = ws.cell(row=row, column=url_col).value
            ws.cell(row=row, column=url_col).style = "Hyperlink"

        if not primary:
            no_image += 1
        elif image_loader is not None:
            try:
                raw_image = image_loader(primary)
                normalised = _normalise_thumbnail(raw_image) if raw_image else None
                if normalised:
                    img = XLImage(io.BytesIO(normalised))
                    img.width = 96
                    img.height = 96
                    ws.add_image(img, f"A{row}")
                    embedded += 1
                else:
                    image_failures += 1
            except Exception:
                image_failures += 1

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 1)}"
    if ws.max_row > 1:
        table = Table(displayName="ProductCatalogue", ref=f"A1:{get_column_letter(len(headers))}{ws.max_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=False,
            showColumnStripes=False,
        )
        ws.add_table(table)

    widths = {
        "Image": 16,
        "SKU": 18,
        "Product Name": 42,
        "Category": 22,
        "Status": 14,
        "Price Display": 20,
        "Price": 15,
        "Compare At Price": 17,
        "Currency": 11,
        "Stock": 10,
        "Featured": 11,
        "Badge": 18,
        "Short Description": 48,
        "Description": 72,
        "Tags": 42,
        "Primary Image URL": 48,
        "All Image URLs": 60,
        "Product URL": 50,
        "Product ID": 38,
        "Created At": 24,
        "Updated At": 24,
    }
    for idx, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(header, 22 if header.startswith("Spec: ") else 18)

    summary = wb.create_sheet("Summary")
    summary.sheet_view.showGridLines = False
    summary["A1"] = "Samrat Glass Emporium — Full Product Catalogue Export"
    summary["A1"].font = Font(size=18, bold=True, color=dark)
    summary.merge_cells("A1:D1")
    summary.append(["Generated (UTC)", generated_at.isoformat(), "", ""])
    summary.append(["Total products", len(ordered), "", ""])
    summary.append(["Embedded primary images", embedded, "", ""])
    summary.append(["Products without an image", no_image, "", ""])
    summary.append(["Images not embedded", image_failures, "", ""])
    summary.append([])
    summary.append(["Status", "Count", "Category", "Count"])
    status_rows = sorted(status_counts.items(), key=lambda x: x[0].casefold())
    category_rows = sorted(category_counts.items(), key=lambda x: x[0].casefold())
    for i in range(max(len(status_rows), len(category_rows))):
        status_item = status_rows[i] if i < len(status_rows) else ("", "")
        category_item = category_rows[i] if i < len(category_rows) else ("", "")
        summary.append([status_item[0], status_item[1], category_item[0], category_item[1]])
    for cell in summary[8]:
        cell.fill = PatternFill("solid", fgColor=dark)
        cell.font = Font(color="FFFFFF", bold=True)
    for col, width in {"A": 28, "B": 18, "C": 30, "D": 18}.items():
        summary.column_dimensions[col].width = width

    output = io.BytesIO()
    wb.save(output)
    metadata = {
        "total": len(ordered),
        "embedded_images": embedded,
        "image_failures": image_failures,
        "no_image": no_image,
    }
    return output.getvalue(), metadata


def _server_workbook(server_module, products: list[dict]) -> tuple[bytes, dict]:
    get_object = server_module.get_object
    return build_catalogue_workbook(
        products,
        image_loader=lambda url: _primary_image_bytes(url, get_object),
    )


def install_catalogue_excel(load_admin_func) -> None:
    """Install GET /api/admin/catalogue/products.xlsx on the active app."""
    server_module = _find_server_module()
    if server_module is None:
        return
    app = server_module.app
    if getattr(app.state, "sge_catalogue_excel_installed", False):
        return

    @app.get("/api/admin/catalogue/products.xlsx")
    async def admin_catalogue_products_xlsx(request: Request):
        user = await load_admin_func(server_module.db, request)
        if user is None:
            if request.cookies.get("session_token") or request.headers.get("Authorization"):
                raise HTTPException(status_code=403, detail="Not authorized for admin.")
            raise HTTPException(status_code=401, detail="Authentication required.")

        products = await server_module.db.products.find(
            {}, {"_id": 0}
        ).sort("sku", 1).to_list(length=10000)
        payload, metadata = await asyncio.to_thread(_server_workbook, server_module, products)
        stamp = datetime.now(timezone.utc).strftime("%Y-%m-%d")
        return Response(
            content=payload,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Cache-Control": "private, no-store",
                "Content-Disposition": f'attachment; filename="samrat-glass-full-product-catalogue-{stamp}.xlsx"',
                "X-Catalogue-Products": str(metadata["total"]),
                "X-Catalogue-Embedded-Images": str(metadata["embedded_images"]),
                "X-Catalogue-Image-Failures": str(metadata["image_failures"]),
            },
        )

    app.state.sge_catalogue_excel_installed = True
