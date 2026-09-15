"""Guarded Admin re-import for catalogue workbooks exported by the site.

Only non-empty ``Spec: ...`` cells are writable. Products are matched by SKU
and cross-checked against the immutable Product ID exported beside each row.
Workbook formatting, filters, hidden rows, images and every non-spec column are
ignored. Preview and apply both rebuild the same deterministic plan so a stale
preview can never be confirmed.
"""
from __future__ import annotations

import hashlib
import io
import json
import zipfile
from dataclasses import dataclass
from typing import Any

from fastapi import File, Form, HTTPException, Request, UploadFile
from openpyxl import load_workbook


MAX_UPLOAD_BYTES = 50 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_ZIP_ENTRIES = 5000
MAX_ROWS = 10000
MAX_COLUMNS = 500
MAX_SPEC_LENGTH = 5000


class CatalogueImportError(ValueError):
    """Raised when an uploaded workbook is not safe or compatible."""


@dataclass(frozen=True)
class ImportRow:
    row_number: int
    sku: str
    product_id: str
    specs: dict[str, str]


@dataclass(frozen=True)
class HistoryAdmin:
    user_id: str
    email: str


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _validate_xlsx_container(payload: bytes) -> None:
    if not payload or len(payload) > MAX_UPLOAD_BYTES:
        raise CatalogueImportError("The Excel file is empty or larger than 50 MB.")
    if payload[:2] != b"PK":
        raise CatalogueImportError("Upload an .xlsx file downloaded from Catalogue Admin.")
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                raise CatalogueImportError("The Excel file contains too many internal files.")
            if sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_BYTES:
                raise CatalogueImportError("The expanded Excel file is too large to import safely.")
    except zipfile.BadZipFile as exc:
        raise CatalogueImportError("The uploaded file is not a valid .xlsx workbook.") from exc


def parse_catalogue_workbook(payload: bytes) -> list[ImportRow]:
    """Read every data row from Products, including rows hidden by Excel filters."""
    _validate_xlsx_container(payload)
    try:
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=False)
    except Exception as exc:
        raise CatalogueImportError("The Excel workbook could not be opened.") from exc
    try:
        if "Products" not in workbook.sheetnames:
            raise CatalogueImportError("The workbook must contain the exported Products sheet.")
        sheet = workbook["Products"]
        if sheet.max_row > MAX_ROWS + 1 or sheet.max_column > MAX_COLUMNS:
            raise CatalogueImportError("The Products sheet is larger than the supported catalogue format.")

        header_cells = next(sheet.iter_rows(min_row=1, max_row=1), ())
        headers = [_cell_text(cell.value) for cell in header_cells]
        normalized_headers = [header.casefold() for header in headers if header]
        if len(normalized_headers) != len(set(normalized_headers)):
            raise CatalogueImportError("The Products sheet contains duplicate column headings.")
        header_index = {header.casefold(): index for index, header in enumerate(headers) if header}
        if "sku" not in header_index or "product id" not in header_index:
            raise CatalogueImportError("The Products sheet must contain SKU and Product ID columns.")

        spec_columns: list[tuple[int, str]] = []
        seen_spec_keys: set[str] = set()
        for index, header in enumerate(headers):
            if not header.casefold().startswith("spec:"):
                continue
            key = header.split(":", 1)[1].strip()
            if not key:
                raise CatalogueImportError("A specification column has no field name.")
            normalized_key = key.casefold()
            if normalized_key in seen_spec_keys:
                raise CatalogueImportError(f'Duplicate specification column: "{key}".')
            seen_spec_keys.add(normalized_key)
            spec_columns.append((index, key))
        if not spec_columns:
            raise CatalogueImportError("No editable Spec: columns were found in the Products sheet.")

        sku_index = header_index["sku"]
        product_id_index = header_index["product id"]
        rows: list[ImportRow] = []
        seen_skus: dict[str, int] = {}
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
            values = [cell.value for cell in cells]
            if not any(value not in (None, "") for value in values):
                continue
            sku = _cell_text(values[sku_index] if sku_index < len(values) else "")
            product_id = _cell_text(values[product_id_index] if product_id_index < len(values) else "")
            if not sku:
                raise CatalogueImportError(f"Row {row_number} has product data but no SKU.")
            normalized_sku = sku.casefold()
            if normalized_sku in seen_skus:
                raise CatalogueImportError(
                    f'SKU "{sku}" appears more than once (rows {seen_skus[normalized_sku]} and {row_number}).'
                )
            seen_skus[normalized_sku] = row_number
            if not product_id:
                raise CatalogueImportError(f'Row {row_number} ({sku}) has no Product ID.')

            specs: dict[str, str] = {}
            for index, key in spec_columns:
                if index >= len(cells):
                    continue
                cell = cells[index]
                if cell.data_type == "f":
                    raise CatalogueImportError(
                        f'Row {row_number} ({sku}) contains a formula in "Spec: {key}". Replace it with a value.'
                    )
                value = _cell_text(cell.value)
                if not value:  # Blank means leave the Admin value unchanged.
                    continue
                if len(value) > MAX_SPEC_LENGTH:
                    raise CatalogueImportError(f'Row {row_number} ({sku}) has an excessively long "Spec: {key}" value.')
                specs[key] = value
            rows.append(ImportRow(row_number, sku, product_id, specs))
        if not rows:
            raise CatalogueImportError("The Products sheet contains no product rows.")
        return rows
    finally:
        workbook.close()


def build_import_plan(rows: list[ImportRow], products: list[dict]) -> dict:
    """Create a deterministic SKU-matched preview without mutating products."""
    products_by_sku: dict[str, dict] = {}
    duplicate_database_skus: set[str] = set()
    for product in products:
        key = _cell_text(product.get("sku")).casefold()
        if not key:
            continue
        if key in products_by_sku:
            duplicate_database_skus.add(key)
        products_by_sku[key] = product

    errors: list[dict] = []
    plan: list[dict] = []
    matched = 0
    for row in rows:
        sku_key = row.sku.casefold()
        product = products_by_sku.get(sku_key)
        if sku_key in duplicate_database_skus:
            errors.append({"row": row.row_number, "sku": row.sku, "message": "This SKU is duplicated in Admin."})
            continue
        if product is None:
            errors.append({"row": row.row_number, "sku": row.sku, "message": "SKU does not exist in Admin."})
            continue
        if _cell_text(product.get("id")) != row.product_id:
            errors.append({
                "row": row.row_number,
                "sku": row.sku,
                "message": "SKU and Product ID do not match. The SKU may have been changed in Excel.",
            })
            continue
        matched += 1
        current_specs = product.get("specs") if isinstance(product.get("specs"), dict) else {}
        updated_specs = dict(current_specs)
        changes = []
        for key, new_value in row.specs.items():
            old_value = _cell_text(current_specs.get(key))
            if old_value == new_value:
                continue
            updated_specs[key] = new_value
            changes.append({"field": f"Spec: {key}", "old": old_value, "new": new_value})
        if changes:
            plan.append({
                "id": product["id"],
                "sku": row.sku,
                "name": _cell_text(product.get("name")),
                "row": row.row_number,
                "changes": changes,
                "before_specs": dict(current_specs),
                "specs": updated_specs,
            })

    token_rows = [
        {
            "id": item["id"],
            "sku": item["sku"],
            "changes": item["changes"],
        }
        for item in sorted(plan, key=lambda item: (item["sku"].casefold(), item["id"]))
    ]
    preview_token = hashlib.sha256(
        json.dumps(token_rows, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    return {
        "workbook_rows": len(rows),
        "matched_count": matched,
        "changed_product_count": len(plan),
        "changed_field_count": sum(len(item["changes"]) for item in plan),
        "unchanged_count": matched - len(plan),
        "errors": errors,
        "can_apply": bool(plan) and not errors,
        "preview_token": preview_token,
        "items": plan,
    }


async def _read_upload(upload: UploadFile) -> bytes:
    filename = str(upload.filename or "")
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(422, "Upload an .xlsx catalogue file.")
    payload = await upload.read(MAX_UPLOAD_BYTES + 1)
    if len(payload) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "The Excel file is larger than 50 MB.")
    return payload


async def _prepare(server_module, payload: bytes) -> dict:
    try:
        rows = parse_catalogue_workbook(payload)
    except CatalogueImportError as exc:
        raise HTTPException(422, str(exc)) from exc
    skus = list(dict.fromkeys(row.sku for row in rows))
    products = await server_module.db.products.find(
        {"sku": {"$in": skus}}, {"_id": 0}
    ).to_list(len(skus) + 100)
    return build_import_plan(rows, products)


def _public_preview(plan: dict) -> dict:
    return {key: value for key, value in plan.items() if key != "items"} | {
        "items": [
            {key: value for key, value in item.items() if key not in {"specs", "before_specs"}}
            for item in plan["items"]
        ]
    }


def history_admin(admin: Any) -> HistoryAdmin:
    """Normalize auth.load_admin's mapping for server history helpers."""
    if isinstance(admin, dict):
        user_id = _cell_text(admin.get("user_id"))
        email = _cell_text(admin.get("email"))
    else:
        user_id = _cell_text(getattr(admin, "user_id", ""))
        email = _cell_text(getattr(admin, "email", ""))
    if not user_id or not email:
        raise HTTPException(403, "The administrator session is missing audit identity details.")
    return HistoryAdmin(user_id=user_id, email=email)


def selected_import_plan(plan: dict, raw_selection: str) -> list[dict]:
    """Return only explicitly selected, already-previewed specification changes."""
    try:
        requested = json.loads(raw_selection)
    except (TypeError, json.JSONDecodeError) as exc:
        raise HTTPException(422, "The selected Excel changes are invalid. Generate a fresh preview.") from exc
    if not isinstance(requested, list) or len(requested) > MAX_ROWS:
        raise HTTPException(422, "The selected Excel changes are invalid. Generate a fresh preview.")

    plan_by_id = {item["id"]: item for item in plan["items"]}
    selected_by_id: dict[str, set[str]] = {}
    for entry in requested:
        if not isinstance(entry, dict):
            raise HTTPException(422, "The selected Excel changes are invalid. Generate a fresh preview.")
        product_id = _cell_text(entry.get("id"))
        fields = entry.get("fields")
        if product_id not in plan_by_id or not isinstance(fields, list):
            raise HTTPException(422, "A selected product is not present in the current preview.")
        allowed = {change["field"] for change in plan_by_id[product_id]["changes"]}
        clean_fields = {_cell_text(field) for field in fields if _cell_text(field)}
        if not clean_fields.issubset(allowed):
            raise HTTPException(422, "A selected specification is not present in the current preview.")
        selected_by_id.setdefault(product_id, set()).update(clean_fields)

    selected_items = []
    for item in plan["items"]:
        selected_fields = selected_by_id.get(item["id"], set())
        changes = [change for change in item["changes"] if change["field"] in selected_fields]
        if not changes:
            continue
        specs = dict(item["before_specs"])
        for change in changes:
            key = change["field"].split("Spec:", 1)[1].strip()
            specs[key] = change["new"]
        selected_items.append({**item, "changes": changes, "specs": specs})
    if not selected_items:
        raise HTTPException(422, "Select at least one specification change to apply.")
    return selected_items


async def rollback_import(server_module, updated_items: list[dict], version_ids: list[str]) -> None:
    """Best-effort compensation if an unexpected write fails mid-import."""
    for record in reversed(updated_items):
        update = {"$set": {"specs": record["before_specs"]}}
        if record["had_updated_at"]:
            update["$set"]["updated_at"] = record["before_updated_at"]
        else:
            update["$unset"] = {"updated_at": ""}
        await server_module.db.products.update_one(
            {"id": record["id"], "specs": record["applied_specs"], "updated_at": record["applied_at"]},
            update,
        )
    if version_ids:
        await server_module.db.product_versions.delete_many({"id": {"$in": version_ids}})


def install_catalogue_excel_import(load_admin_func, require_csrf_func) -> None:
    """Install guarded preview/apply endpoints on the active FastAPI app."""
    from catalogue_excel import _find_server_module

    server_module = _find_server_module()
    if server_module is None:
        return
    app = server_module.app
    if getattr(app.state, "sge_catalogue_excel_import_installed", False):
        return

    async def require_admin_upload(request: Request):
        require_csrf_func(request)
        user = await load_admin_func(server_module.db, request)
        if user is None:
            if request.cookies.get("session_token") or request.headers.get("Authorization"):
                raise HTTPException(403, "Not authorized for admin.")
            raise HTTPException(401, "Authentication required.")
        return user

    @app.post("/api/admin/catalogue/import/preview")
    async def preview_catalogue_import(request: Request, file: UploadFile = File(...)):
        await require_admin_upload(request)
        plan = await _prepare(server_module, await _read_upload(file))
        return _public_preview(plan)

    @app.post("/api/admin/catalogue/import/apply")
    async def apply_catalogue_import(
        request: Request,
        file: UploadFile = File(...),
        preview_token: str = Form(...),
        reason: str = Form(...),
        selected_changes: str = Form(...),
    ):
        admin = await require_admin_upload(request)
        clean_reason = str(reason or "").strip()
        if not clean_reason:
            raise HTTPException(422, "Enter a reason for the catalogue import.")
        if len(clean_reason) > 500:
            raise HTTPException(422, "The import reason must be 500 characters or fewer.")
        plan = await _prepare(server_module, await _read_upload(file))
        if plan["errors"]:
            raise HTTPException(422, "Resolve every workbook error and generate a fresh preview.")
        if not plan["items"]:
            raise HTTPException(422, "The workbook contains no specification changes.")
        if plan["preview_token"] != str(preview_token or ""):
            raise HTTPException(409, "Catalogue data changed after preview. Generate a fresh preview before applying.")

        selected_items = selected_import_plan(plan, selected_changes)
        audit_admin = history_admin(admin)
        prepared = []
        for item in selected_items:
            existing = await server_module.db.products.find_one({"id": item["id"]}, {"_id": 0})
            if not existing:
                raise HTTPException(409, "A product changed after preview. Generate a fresh preview.")
            existing_specs = existing.get("specs") if isinstance(existing.get("specs"), dict) else {}
            if existing_specs != item["before_specs"]:
                raise HTTPException(409, "A product changed after preview. Generate a fresh preview.")
            prepared.append((item, existing))

        changed_ids = []
        updated_items = []
        version_ids = []
        try:
            # Create every baseline before the first catalogue mutation.
            for _, existing in prepared:
                await server_module._ensure_product_history_baseline(existing, audit_admin)
            for item, existing in prepared:
                guard = {"id": item["id"]}
                if "specs" in existing:
                    guard["specs"] = item["before_specs"]
                else:
                    guard["specs"] = {"$exists": False}
                applied_at = server_module.now_iso()
                result = await server_module.db.products.update_one(
                    guard,
                    {"$set": {"specs": item["specs"], "updated_at": applied_at}},
                )
                if result.modified_count != 1:
                    raise HTTPException(409, "A product changed while applying the import. Generate a fresh preview.")
                updated_items.append({
                    "id": item["id"],
                    "before_specs": item["before_specs"],
                    "applied_specs": item["specs"],
                    "applied_at": applied_at,
                    "had_updated_at": "updated_at" in existing,
                    "before_updated_at": existing.get("updated_at"),
                })
                updated = await server_module.db.products.find_one({"id": item["id"]}, {"_id": 0})
                version = await server_module._record_product_version(
                    updated,
                    audit_admin,
                    action="excel_import",
                    reason=clean_reason,
                    previous=existing,
                )
                version_ids.append(version["id"])
                changed_ids.append(item["id"])
        except Exception as exc:
            try:
                await rollback_import(server_module, updated_items, version_ids)
            except Exception:
                server_module.logger.exception("Excel import rollback failed")
            if isinstance(exc, HTTPException):
                raise
            server_module.logger.exception("Excel catalogue import failed")
            raise HTTPException(500, "The import could not be applied. Any completed product writes were rolled back.") from exc
        return {
            "ok": True,
            "updated_count": len(changed_ids),
            "changed_field_count": sum(len(item["changes"]) for item in selected_items),
            "ids": changed_ids,
        }

    app.state.sge_catalogue_excel_import_installed = True
